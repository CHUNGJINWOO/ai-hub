from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def extract_text_from_docx(
    path: Path,
) -> list[tuple[int | None, str]]:
    doc = Document(path)
    sections = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        rows = []

        for row in table.rows:
            values = [_clean(cell.text) for cell in row.cells]
            rows.append(" | ".join(values))

        table_text = "\n".join(line for line in rows if line.strip())

        if table_text:
            sections.append(
                f"[Table {table_index}]\n{table_text}"
            )

    text = "\n\n".join(sections).strip()

    if not text:
        return []

    return [(None, text)]


def extract_text_from_xlsx(
    path: Path,
) -> list[tuple[int | None, str]]:
    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    pages = []

    for sheet in workbook.worksheets:
        rows = []

        for row in sheet.iter_rows(values_only=True):
            values = [_clean(value) for value in row]

            if any(values):
                rows.append(" | ".join(values))

        if rows:
            text = (
                f"[Sheet: {sheet.title}]\n"
                + "\n".join(rows)
            )
            pages.append((None, text))

    workbook.close()

    return pages


def extract_text_from_pptx(
    path: Path,
) -> list[tuple[int | None, str]]:
    presentation = Presentation(path)
    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        parts = []

        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:
                parts.append(text)

        for table_index, shape in enumerate(
            slide.shapes,
            start=1,
        ):
            if not getattr(shape, "has_table", False):
                continue

            rows = []

            for row in shape.table.rows:
                values = [
                    _clean(cell.text)
                    for cell in row.cells
                ]
                rows.append(" | ".join(values))

            table_text = "\n".join(
                line for line in rows if line.strip()
            )

            if table_text:
                parts.append(
                    f"[Table {table_index}]\n{table_text}"
                )

        if parts:
            text = (
                f"[Slide {slide_number}]\n"
                + "\n\n".join(parts)
            )
            slides.append((slide_number, text))

    return slides
